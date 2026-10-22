# agentic-spreadsheet-processor/src/spreadsheet_agent.py
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, IconSetRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter, range_boundaries
import logging
import datetime
import json
from utils import setup_logging

logger = setup_logging(__name__)

class SpreadsheetAgent:
    """An agent that performs operations on an Excel workbook based on configuration."""
    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []
        self.start_time = datetime.datetime.now()
        logger.info("Spreadsheet Agent initialized.")

    def _log_audit(self, status, message, task_type="N/A", details=None):
        log_entry = {
            "timestamp": datetime.datetime.now().isoformat(),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details if details else {}
        }
        self.audit_log.append(log_entry)
        if status in ["ERROR", "CRITICAL"]:
            logger.error(f"AUDIT LOG - {status}: {message} | Task: {task_type} | Details: {details}")
        else:
            logger.info(f"AUDIT LOG - {status}: {message} | Task: {task_type}")

    def _load_workbook(self):
        input_file = self.config["input_file"]
        try:
            self.workbook = openpyxl.load_workbook(input_file)
            self._log_audit("SUCCESS", f"Workbook '{input_file}' loaded.", "load_workbook")
            logger.info(f"Loaded workbook: {input_file}")
            return True
        except FileNotFoundError:
            self._log_audit("CRITICAL", f"Input file not found: {input_file}", "load_workbook")
            return False
        except Exception as e:
            self._log_audit("CRITICAL", f"Failed to load workbook '{input_file}': {e}", "load_workbook")
            return False

    def _get_sheet(self, sheet_name):
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        else:
            self._log_audit("WARNING", f"Sheet '{sheet_name}' not found. Creating new sheet.", "get_sheet")
            return self.workbook.create_sheet(sheet_name)

    def _process_copy_data_task(self, task):
        src_sheet_name = task.get("source_sheet")
        dest_sheet_name = task.get("dest_sheet")
        if not src_sheet_name or not dest_sheet_name:
            self._log_audit("ERROR", "Missing source_sheet or dest_sheet for copy_data task.", "copy_data", task)
            return

        src_sheet = self._get_sheet(src_sheet_name)
        dest_sheet = self._get_sheet(dest_sheet_name)

        for row in src_sheet.iter_rows():
            for cell in row:
                dest_sheet[cell.coordinate].value = cell.value

        self._log_audit("SUCCESS", f"Data copied from '{src_sheet_name}' to '{dest_sheet_name}'.", "copy_data")

    def _process_calculate_column_task(self, task):
        sheet_name = task.get("sheet")
        target_column_letter = task.get("target_column")
        start_row = task.get("start_row", 2)
        formula = task.get("formula")

        sheet = self._get_sheet(sheet_name)
        header_name = task.get("header", "Calculated Column")
        sheet[f"{target_column_letter}1"].value = header_name
        sheet[f"{target_column_letter}1"].font = Font(bold=True)

        for row_num in range(start_row, sheet.max_row + 1):
            cell_coord = f"{target_column_letter}{row_num}"
            try:
                sheet[cell_coord] = formula.replace("{row}", str(row_num))
                if task.get("number_format"):
                    sheet[cell_coord].number_format = task["number_format"]
            except Exception as e:
                self._log_audit("ERROR", f"Failed formula at {cell_coord}: {e}", "calculate_column", task)

        self._log_audit("SUCCESS", f"Column '{target_column_letter}' calculated in '{sheet_name}'.", "calculate_column")

    def _process_conditional_format_task(self, task):
        sheet_name = task.get("sheet")
        data_range = task.get("range")
        rule_type = task.get("rule_type")

        sheet = self._get_sheet(sheet_name)
        try:
            if rule_type == "color_scale":
                rule = ColorScaleRule(
                    start_type='min', start_color=task.get("min_color", "FF0000"),
                    mid_type='percentile', mid_value=50, mid_color=task.get("mid_color", "FFFF00"),
                    end_type='max', end_color=task.get("max_color", "00FF00")
                )
                sheet.conditional_formatting.add(data_range, rule)
            elif rule_type == "expression":
                formula = task.get("formula")
                fill_color = task.get("fill_color", "FFC7CE")
                font_color = task.get("font_color", "9C0006")
                dxf = DifferentialStyle(font=Font(color=font_color), fill=PatternFill(bgColor=fill_color))
                sheet.conditional_formatting.add(data_range, ruleType='expression', formula=[formula], dxf=dxf)
            self._log_audit("SUCCESS", f"Applied conditional format '{rule_type}' to '{data_range}'.", "conditional_format")
        except Exception as e:
            self._log_audit("ERROR", f"Failed conditional formatting: {e}", "conditional_format", task)

    def _process_chart_task(self, task):
        sheet_name = task.get("sheet")
        chart_type = task.get("chart_type")
        data_range = task.get("data_range")
        category_range = task.get("category_range")

        sheet = self._get_sheet(sheet_name)
        try:
            chart = BarChart() if chart_type == "bar" else LineChart() if chart_type == "line" else PieChart()
            min_col, min_row, max_col, max_row = range_boundaries(data_range)
            data = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.add_data(data, titles_from_data=True)

            if category_range:
                min_c, min_r, max_c, max_r = range_boundaries(category_range)
                cats = Reference(sheet, min_col=min_c, min_row=min_r, max_col=max_c, max_row=max_r)
                chart.set_categories(cats)

            chart.title = task.get("title", "Chart")
            sheet.add_chart(chart, task.get("top_left_cell", "G2"))
            self._log_audit("SUCCESS", f"Chart '{chart.title}' added to '{sheet_name}'.", "chart")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to create chart: {e}", "chart", task)

    def _add_audit_log_sheet(self):
        audit_sheet = self.workbook.create_sheet("Audit Log", 0)
        audit_sheet.append(["Timestamp", "Status", "Task Type", "Message", "Details"])
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
        for col in range(1, 6):
            cell = audit_sheet.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill

        for entry in self.audit_log:
            audit_sheet.append([
                entry["timestamp"], entry["status"], entry["task_type"],
                entry["message"], json.dumps(entry["details"])
            ])

    def run(self):
        if not self._load_workbook():
            return

        for task in self.config.get("tasks", []):
            ttype = task.get("type")
            if ttype == "copy_data": self._process_copy_data_task(task)
            elif ttype == "calculate_column": self._process_calculate_column_task(task)
            elif ttype == "conditional_format": self._process_conditional_format_task(task)
            elif ttype == "chart": self._process_chart_task(task)

        self._add_audit_log_sheet()
        output_file = self.config["output_file"]
        try:
            self.workbook.save(output_file)
            self._log_audit("SUCCESS", f"Workbook saved to '{output_file}'.", "save_workbook")
            logger.info(f"Report successfully saved to: {output_file}")
        except Exception as e:
            logger.critical(f"Error saving file: {e}")